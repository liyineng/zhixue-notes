"""Parsing inventory and retries use actual local file content; models are mocked."""
import asyncio
import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import AsyncMock, patch

import pymupdf
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from app import db
from app.document_status import inspect_document, retry_document_pages
from app.document_images import render_region
from app.parsers import parse_document


class DocumentStatusTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.database = patch.object(db, 'DB_PATH', self.root / 'test.db')
        self.database.start()
        db.init_db()

    def tearDown(self):
        self.database.stop()
        self.temp.cleanup()

    def image(self):
        stream = io.BytesIO()
        Image.new('RGB', (1000, 300), 'white').save(stream, format='PNG')
        return stream.getvalue()

    def pdf(self):
        path = self.root / 'course.pdf'
        with pymupdf.open() as pdf:
            first = pdf.new_page(width=600, height=500)
            first.insert_text((20, 40), 'Native source text about circuit analysis and valid formula conditions.')
            second = pdf.new_page(width=600, height=500)
            second.insert_image(pymupdf.Rect(20, 20, 580, 200), stream=self.image())
            pdf.new_page(width=600, height=500)
            fourth = pdf.new_page(width=600, height=500)
            fourth.insert_text((20, 40), 'This page has significant native text plus a separate image-based exercise.')
            fourth.insert_image(pymupdf.Rect(20, 120, 580, 300), stream=self.image())
            pdf.save(path)
        return self.register(path)

    def register(self, path):
        chunks, count = parse_document(path, path.name)
        document = {'id': path.stem, 'name': path.name, 'path': str(path), 'kind': path.suffix[1:],
                    'page_count': count, 'chunk_count': len(chunks), 'size_bytes': path.stat().st_size}
        db.insert_document(document, chunks)
        return document

    def presentation(self):
        path = self.root / 'slides.pptx'
        deck = Presentation()
        for number in range(3):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            text = slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(7), Inches(.6))
            text.text_frame.text = f'Course slide {number + 1}'
            if number == 1:
                slide.shapes.add_picture(io.BytesIO(self.image()), Inches(.5), Inches(1.3), width=Inches(7))
        deck.save(path)
        # Actual embedded formula marker; text anchors alone are not formula evidence.
        with zipfile.ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        entries['ppt/slides/slide3.xml'] = entries['ppt/slides/slide3.xml'].replace(b'</p:spTree>', b'<p:oleObj progId="Equation.3"/></p:spTree>')
        with zipfile.ZipFile(path, 'w') as archive:
            for name, data in entries.items(): archive.writestr(name, data)
        return self.register(path)

    def test_inventory_reads_original_pages_not_citation_gaps_and_never_calls_models(self):
        document = self.pdf()
        with patch('app.document_images.read_question', AsyncMock()) as model:
            status = inspect_document(document)
        model.assert_not_called()
        self.assertEqual(status['page_count'], 4)
        self.assertEqual(status['summary']['native_text_pages'], [1, 4])
        self.assertEqual(status['summary']['indexed_text_pages'], [1, 4])
        self.assertEqual(status['summary']['pending_pages'], [2, 4])
        self.assertEqual(status['summary']['empty_pages'], [3])
        self.assertFalse(status['pages'][2]['retryable'])
        self.assertIn('矢量图', status['pages'][2]['detail'])
        self.assertIn('不代表已覆盖所有知识点', status['notice'])

    def test_pdf_retry_selected_page_preserves_original_chunks_and_caches_success(self):
        document = self.pdf()
        original_bytes = Path(document['path']).read_bytes()
        ids = db.document_chunk_ids(document['id'])
        progress = []
        model = AsyncMock(return_value=(r'Question: find the frequency \(f=1/T\).', 'vision'))
        with patch('app.document_images.read_question', model):
            status = asyncio.run(retry_document_pages(document, {'vision_model': 'mock'}, [2], progress.append))
        self.assertEqual(progress, [2])
        self.assertEqual(model.await_count, 1)
        self.assertEqual(status['retry']['requested_pages'], [2])
        self.assertEqual(status['retry']['added_pages'], [2])
        self.assertEqual(status['retry']['unread_pages'], [])
        self.assertEqual(status['summary']['recognized_image_pages'], [2])
        self.assertEqual(status['summary']['pending_pages'], [4])
        self.assertTrue(ids.issubset(db.document_chunk_ids(document['id'])))
        self.assertEqual(Path(document['path']).read_bytes(), original_bytes)
        with self.assertRaisesRegex(ValueError, '没有待识别'):
            asyncio.run(retry_document_pages(document, {'vision_model': 'mock'}, [2]))
        self.assertEqual(inspect_document(document)['pages'][1]['status'], 'recognized')

    def test_failed_page_stays_retryable_and_uncertain_vision_is_not_marked_read(self):
        document = self.pdf()
        with patch('app.document_images.read_question', AsyncMock(return_value=('[看不清]', 'vision'))), \
             patch('app.document_images.ocr_lines', return_value=[]):
            status = asyncio.run(retry_document_pages(document, {'vision_model': 'mock', 'api_key': 'never-in-report'}, [2]))
        self.assertEqual(status['retry']['unread_pages'], [2])
        self.assertEqual(status['pages'][1]['status'], 'failed')
        self.assertTrue(status['pages'][1]['retryable'])
        self.assertNotIn('image-course-2', db.document_chunk_ids(document['id']))
        records = list((self.root / 'document-status').glob('*.json'))
        self.assertEqual(len(records), 1)
        self.assertNotIn('never-in-report', records[0].read_text())
        with patch('app.document_images.read_question', AsyncMock(return_value=('The complete exercise.', 'vision'))):
            complete = asyncio.run(retry_document_pages(document, {'vision_model': 'mock'}, [2]))
        self.assertEqual(complete['pages'][1]['status'], 'recognized')
        self.assertEqual(complete['summary']['failed_pages'], [])

    def test_empty_invalid_blank_and_text_only_retry_fail_closed(self):
        document = self.pdf()
        for pages in ([], [0], [99], [True], [1], [3]):
            with self.subTest(pages=pages), self.assertRaises(ValueError):
                asyncio.run(retry_document_pages(document, {}, pages))
        self.assertFalse((self.root / 'document-status').exists())

    def test_later_page_failure_keeps_earlier_success_and_signals_index_refresh(self):
        document = self.pdf()
        def render(path, page, bbox):
            if page == 4:
                raise RuntimeError('模拟后续页渲染失败')
            return render_region(path, page, bbox)
        with patch('app.document_images.render_region', side_effect=render), \
             patch('app.document_images.read_question', AsyncMock(return_value=('Complete exercise.', 'vision'))):
            status = asyncio.run(retry_document_pages(document, {'vision_model': 'mock'}, [2, 4]))
        self.assertEqual(status['retry']['added_pages'], [2])
        self.assertEqual(status['retry']['unread_pages'], [4])
        self.assertEqual(status['pages'][1]['status'], 'recognized')
        self.assertEqual(status['pages'][3]['status'], 'failed')

    def test_ppt_images_and_formula_objects_are_separate_and_retry_only_selected(self):
        document = self.presentation()
        status = inspect_document(document)
        self.assertEqual(status['summary']['native_text_pages'], [1, 2, 3])
        self.assertEqual(status['summary']['pending_pages'], [2, 3])
        self.assertTrue(status['pages'][1]['image_candidate'])
        self.assertFalse(status['pages'][1]['formula_candidate'])
        self.assertTrue(status['pages'][2]['formula_candidate'])
        preview = self.root / 'preview.pdf'
        with pymupdf.open() as pdf:
            for _ in range(3): pdf.new_page(width=600, height=500)
            pdf.save(preview)
        ids = db.document_chunk_ids(document['id'])
        progress = []
        with patch('app.document_images.presentation_pdf', return_value=preview), \
             patch('app.document_images.read_question', AsyncMock(return_value=(r'\[Q=\omega L/R\]', 'vision'))) as model:
            result = asyncio.run(retry_document_pages(document, {'vision_model': 'mock'}, [3], progress.append))
        self.assertEqual(progress, [3])
        self.assertEqual(model.await_count, 1)
        self.assertEqual(result['summary']['recognized_formula_pages'], [3])
        self.assertEqual(result['summary']['pending_pages'], [2])
        with patch('app.document_images.presentation_pdf', return_value=preview), \
             patch('app.document_images.read_question', AsyncMock(return_value=('Complete exercise and diagram relations.', 'vision'))):
            result = asyncio.run(retry_document_pages(document, {'vision_model': 'mock'}, [2]))
        self.assertEqual(result['summary']['recognized_image_pages'], [2])
        self.assertEqual(result['summary']['pending_pages'], [])
        self.assertTrue(ids.issubset(db.document_chunk_ids(document['id'])))

    def test_missing_original_is_explicit_and_does_not_destroy_cached_text(self):
        document = self.pdf()
        ids = db.document_chunk_ids(document['id'])
        Path(document['path']).unlink()
        result = inspect_document(document)
        self.assertFalse(result['available'])
        self.assertIn('原文件缺失', result['error'])
        self.assertEqual(ids, db.document_chunk_ids(document['id']))
        with self.assertRaises(ValueError): asyncio.run(retry_document_pages(document, {}, [2]))
