import io
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from fastapi.testclient import TestClient
from pptx import Presentation

from app import db, main, previews


class PresentationPreviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.db_patch = patch.object(db, 'DB_PATH', self.root / 'app.db')
        self.upload_patch = patch.object(main, 'UPLOAD_DIR', self.root / 'uploads')
        self.db_patch.start()
        self.upload_patch.start()
        main.UPLOAD_DIR.mkdir()
        self.client = TestClient(main.app)
        self.client.__enter__()
        deck = Presentation()
        for text in ['First slide', 'Second slide']:
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = text
        stream = io.BytesIO()
        deck.save(stream)
        self.content = stream.getvalue()
        response = self.client.post('/api/documents', files={'file': ('slides.pptx', self.content)})
        self.assertEqual(response.status_code, 200, response.text)
        self.document = response.json()['document']

    def tearDown(self):
        self.client.__exit__(None, None, None)
        self.upload_patch.stop()
        self.db_patch.stop()
        main._invalidate_index()
        self.temp.cleanup()

    def convert(self, command, **kwargs):
        work = Path(command[command.index('--outdir') + 1])
        self.assertIn('--headless', command)
        self.assertTrue(any(v.startswith('-env:UserInstallation=') for v in command))
        with pymupdf.open() as pdf:
            pdf.new_page()
            pdf.new_page()
            pdf.save(work / (Path(command[-1]).stem + '.pdf'))
        return subprocess.CompletedProcess(command, 0)

    def test_cached_preview_retains_original_chunks_ids_and_page_count(self):
        original_chunks = db.document_chunk_ids(self.document['id'])
        with patch.object(previews, 'office_executable', return_value='office'), patch.object(previews.subprocess, 'run', side_effect=self.convert) as converter:
            for _ in range(2):
                response = self.client.get('/api/documents/' + self.document['id'] + '/preview')
                self.assertEqual(response.status_code, 200, response.text[:200])
                self.assertIn('application/pdf', response.headers['content-type'])
                with pymupdf.open(stream=response.content, filetype='pdf') as pdf:
                    self.assertEqual(pdf.page_count, 2)
            self.assertEqual(converter.call_count, 1)
        self.assertEqual(Path(self.document['path']).read_bytes(), self.content)
        self.assertEqual(db.document_chunk_ids(self.document['id']), original_chunks)
        self.assertEqual(self.client.get('/api/documents/' + self.document['id'] + '/file').content, self.content)

    def test_missing_converter_and_timeout_are_actionable_without_losing_text(self):
        with patch.object(previews, 'office_executable', return_value=None):
            response = self.client.get('/api/documents/' + self.document['id'] + '/preview')
            self.assertEqual(response.status_code, 503)
            self.assertIn('LibreOffice', response.json()['detail'])
        with patch.object(previews, 'office_executable', return_value='office'), patch.object(previews.subprocess, 'run', side_effect=subprocess.TimeoutExpired('office', 120)):
            response = self.client.get('/api/documents/' + self.document['id'] + '/preview')
            self.assertEqual(response.status_code, 503)
            self.assertIn('超时', response.json()['detail'])
        self.assertTrue(self.client.get('/api/documents/' + self.document['id'] + '/chunks').json()['chunks'])
        self.assertFalse(list((self.root / 'previews').glob('convert-*')))

    def test_incomplete_conversion_is_not_used_as_preview(self):
        def one_page(command, **kwargs):
            work = Path(command[command.index('--outdir') + 1])
            with pymupdf.open() as pdf:
                pdf.new_page()
                pdf.save(work / (Path(command[-1]).stem + '.pdf'))
        with patch.object(previews, 'office_executable', return_value='office'), patch.object(previews.subprocess, 'run', side_effect=one_page):
            response = self.client.get('/api/documents/' + self.document['id'] + '/preview')
            self.assertEqual(response.status_code, 503)
            self.assertIn('页数不完整', response.json()['detail'])
        self.assertFalse(list((self.root / 'previews').glob('*.pdf')))
        self.assertEqual(self.client.get('/api/documents/missing/preview').status_code, 404)

    def test_mac_font_environment_and_windows_environment_remain_separate(self):
        import xml.etree.ElementTree as xml
        with patch.object(previews.sys, 'platform', 'win32'), patch.dict(previews.os.environ, {'FONTCONFIG_FILE': 'custom.conf'}):
            self.assertEqual(previews.conversion_environment(self.root)['FONTCONFIG_FILE'], 'custom.conf')
        with patch.object(previews.sys, 'platform', 'darwin'), patch.dict(previews.os.environ, {}, clear=True), patch.object(Path, 'is_file', return_value=False):
            env = previews.conversion_environment(self.root)
            self.assertEqual(Path(env['FONTCONFIG_FILE']), self.root / 'fonts.conf')
            self.assertEqual(xml.parse(env['FONTCONFIG_FILE']).getroot().tag, 'fontconfig')
