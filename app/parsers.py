from __future__ import annotations

import re
from pathlib import Path
from typing import Any
from uuid import uuid4

from .config import CHUNK_OVERLAP, CHUNK_SIZE


class UnsupportedDocument(ValueError):
    pass


def _clean(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _make_chunks(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    chunks: list[dict[str, Any]] = []
    current: list[dict[str, Any]] = []
    current_length = 0

    def flush(items: list[dict[str, Any]]) -> None:
        if not items:
            return
        text = _clean("\n".join(item["text"] for item in items))
        if not text:
            return
        pages = [item.get("page") for item in items if item.get("page") is not None]
        chunks.append(
            {
                "id": uuid4().hex,
                "ordinal": len(chunks),
                "page": min(pages) if pages else None,
                "text": text,
                "lines": [
                    {
                        "page": item.get("page"),
                        "text": item["text"],
                        "bbox": item.get("bbox"),
                    }
                    for item in items
                ],
            }
        )

    for record in records:
        text = _clean(record.get("text", ""))
        if not text:
            continue
        record = {**record, "text": text}
        record_length = len(text) + 1
        if current and current_length + record_length > CHUNK_SIZE:
            flush(current)
            overlap: list[dict[str, Any]] = []
            overlap_length = 0
            for previous in reversed(current):
                candidate = len(previous["text"]) + 1
                if overlap and overlap_length + candidate > CHUNK_OVERLAP:
                    break
                overlap.insert(0, previous)
                overlap_length += candidate
            current = overlap
            current_length = overlap_length
        current.append(record)
        current_length += record_length

    flush(current)
    return chunks


def _parse_pdf(path: Path) -> tuple[list[dict[str, Any]], int]:
    try:
        import pymupdf as fitz
    except ImportError as exc:
        raise RuntimeError("PyMuPDF is required to parse PDF files.") from exc

    records: list[dict[str, Any]] = []
    with fitz.open(path) as document:
        for page_number, page in enumerate(document, start=1):
            data = page.get_text("dict")
            for block in data.get("blocks", []):
                if block.get("type") != 0:
                    continue
                for line in block.get("lines", []):
                    text = "".join(
                        span.get("text", "") for span in line.get("spans", [])
                    )
                    text = _clean(text)
                    if not text:
                        continue
                    bbox = line.get("bbox")
                    records.append(
                        {
                            "page": page_number,
                            "text": text,
                            "bbox": [float(value) for value in bbox] if bbox else None,
                        }
                    )
        return _make_chunks(records), document.page_count


def _parse_pptx(path: Path) -> tuple[list[dict[str, Any]], int]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("python-pptx is required to parse PPTX files.") from exc

    records: list[dict[str, Any]] = []
    presentation = Presentation(str(path))
    for slide_number, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            text_frame = getattr(shape, "text_frame", None)
            if text_frame is None:
                continue
            for paragraph in text_frame.paragraphs:
                text = _clean(paragraph.text)
                if text:
                    records.append(
                        {"page": slide_number, "text": text, "bbox": None}
                    )
    return _make_chunks(records), len(presentation.slides)


def _parse_docx(path: Path) -> tuple[list[dict[str, Any]], int]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("python-docx is required to parse DOCX files.") from exc

    records: list[dict[str, Any]] = []
    document = Document(str(path))
    for paragraph in document.paragraphs:
        text = _clean(paragraph.text)
        if text:
            records.append({"page": None, "text": text, "bbox": None})
    return _make_chunks(records), 1


def _parse_text(path: Path) -> tuple[list[dict[str, Any]], int]:
    text = path.read_text(encoding="utf-8", errors="replace")
    records: list[dict[str, Any]] = []
    section = 1
    for block in re.split(r"\n\s*\n", text):
        block = _clean(block)
        if not block:
            continue
        if block.startswith("#"):
            section += 1
        records.append({"page": section, "text": block, "bbox": None})
    return _make_chunks(records), section


def parse_document(path: Path, original_name: str) -> tuple[list[dict[str, Any]], int]:
    suffix = Path(original_name).suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix in {".md", ".markdown", ".txt"}:
        return _parse_text(path)
    raise UnsupportedDocument(f"Unsupported file type: {suffix}")
