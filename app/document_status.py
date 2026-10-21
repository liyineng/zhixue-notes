"""Read-only source parsing inventory and explicit, page-scoped recognition retries.

The report describes detected content and cached extraction, not knowledge coverage.
Recognition attempt metadata is local, separate from the live SQLite schema.
"""
from __future__ import annotations

import asyncio
import hashlib
import json
import os
import tempfile
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from collections.abc import Callable

import pymupdf
from starlette.concurrency import run_in_threadpool

from . import db
from .document_images import (ensure_image_text, ensure_presentation_formulas, image_regions,
                              presentation_equation_pages, presentation_image_pages)

_retry_locks: dict[str, asyncio.Lock] = {}
_active_retries: set[str] = set()


def _report_path(document: dict[str, Any]) -> Path:
    identity = hashlib.sha256(document['id'].encode()).hexdigest()
    return Path(db.DB_PATH).parent / 'document-status' / f'{identity}.json'


def _fingerprint(document: dict[str, Any]) -> str:
    try:
        stat = Path(document['path']).stat()
        return f'{stat.st_size}:{stat.st_mtime_ns}'
    except OSError:
        return ''


def _read_report(document: dict[str, Any]) -> dict[str, Any]:
    try:
        report = json.loads(_report_path(document).read_text(encoding='utf-8'))
        if isinstance(report, dict) and report.get('fingerprint') == _fingerprint(document) and isinstance(report.get('pages'), dict):
            return report
    except (OSError, ValueError, TypeError):
        pass
    return {'fingerprint': _fingerprint(document), 'pages': {}}


def _write_report(document: dict[str, Any], report: dict[str, Any]) -> None:
    path = _report_path(document)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Atomic publication prevents a stopped recognition request from corrupting metadata.
    with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=path.parent, delete=False) as stream:
        temporary = Path(stream.name)
        try:
            json.dump(report, stream, ensure_ascii=False)
            stream.flush()
            os.fsync(stream.fileno())
        except BaseException:
            temporary.unlink(missing_ok=True)
            raise
    try:
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _chunks(document_id: str) -> list[dict[str, Any]]:
    with db.connection() as conn:
        rows = conn.execute('SELECT id, page, text, lines_json FROM chunks WHERE document_id=?', (document_id,)).fetchall()
    result = []
    for row in rows:
        item = dict(row)
        try:
            item['lines'] = json.loads(item.pop('lines_json'))
            if not isinstance(item['lines'], list):
                item['lines'] = []
        except (ValueError, TypeError):
            item['lines'] = []
        result.append(item)
    return result


def _native_pages(document: dict[str, Any]) -> tuple[int, set[int], set[int], set[int]]:
    """Return actual page count, native text pages, significant image and formula candidates."""
    path = Path(document['path'])
    if document['kind'] == 'pdf':
        with pymupdf.open(path) as pdf:
            count = len(pdf)
            text = {number for number, page in enumerate(pdf, 1) if page.get_text().strip()}
        images = {candidate['page'] for candidate in image_regions(path, include_mixed=True)}
        return count, text, images, set()
    if document['kind'] == 'pptx':
        from pptx import Presentation
        deck = Presentation(str(path))
        def text_in(shapes):
            return any((getattr(shape, 'has_text_frame', False) and shape.text.strip())
                       or (getattr(shape, 'has_table', False) and any(cell.text.strip() for row in shape.table.rows for cell in row.cells))
                       or (hasattr(shape, 'shapes') and text_in(shape.shapes)) for shape in shapes)
        text = {number for number, slide in enumerate(deck.slides, 1) if text_in(slide.shapes)}
        return len(deck.slides), text, set(presentation_image_pages(path)), set(presentation_equation_pages(path))
    return int(document.get('page_count') or 1), set(), set(), set()


def inspect_document(document: dict[str, Any]) -> dict[str, Any]:
    """Inspect originals and existing chunks without OCR, conversion, or model requests."""
    response: dict[str, Any] = {'document_id': document['id'], 'name': document['name'], 'kind': document['kind'],
                                'available': Path(document['path']).is_file(), 'pages': [], 'summary': {}, 'error': ''}
    response['notice'] = ('根据原文件及已缓存的提取结果统计。图片候选也可能是图表；已识别不代表已覆盖所有知识点。'
                          'PDF 中没有独立对象的公式无法自动逐项核验，请结合原文检查。')
    if not response['available']:
        response['error'] = '原文件缺失，保留的索引仍可检索；无法检查或重试页面。'
        return response
    try:
        count, native, images, formulas = _native_pages(document)
    except (OSError, ValueError, RuntimeError, KeyError, IndexError, EOFError, zipfile.BadZipFile, SyntaxError) as error:
        response['error'] = '无法检查原文件的页面结构，请先确认原文件可正常打开。'
        response['error_type'] = type(error).__name__
        return response
    chunks = _chunks(document['id'])
    indexed: set[int] = set()
    recognized: dict[int, set[str]] = {}
    image_cache: set[int] = set()
    formula_cache: set[int] = set()
    for chunk in chunks:
        lines = chunk.get('lines') or []
        for line in lines:
            page = line.get('page') or chunk.get('page')
            if not isinstance(page, int) or not line.get('text', '').strip():
                continue
            method = line.get('image_recognition')
            if method:
                recognized.setdefault(page, set()).add(str(method))
                if not chunk['id'].startswith(f"formulas-v1-{document['id']}-"):
                    image_cache.add(page)
            else:
                indexed.add(page)
        page = chunk.get('page')
        if isinstance(page, int) and chunk.get('text', '').strip():
            if chunk['id'].startswith(f"formulas-v1-{document['id']}-"):
                formula_cache.add(page)
            elif chunk['id'].startswith(f"image-{document['id']}-"):
                image_cache.add(page)
            elif not lines:
                indexed.add(page)
    if document['kind'] not in {'pdf', 'pptx'}:
        native = indexed
        response['notice'] = '文本文档按导入时的章节或段落编号统计，不代表打印页码；本面板不检测其中的内嵌图片。'
    report = _read_report(document)
    for number in range(1, count + 1):
        image_pending = number in images and number not in image_cache
        formula_pending = number in formulas and number not in formula_cache
        pending = image_pending or formula_pending
        last = report['pages'].get(str(number), {})
        if not isinstance(last, dict):
            last = {}
        interrupted = last.get('status') == 'processing' and document['id'] not in _active_retries
        failed = pending and (last.get('status') in {'failed', 'interrupted'} or interrupted)
        processing = pending and document['id'] in _active_retries and last.get('status') == 'processing'
        status = 'processing' if processing else 'failed' if failed else 'pending' if pending else 'recognized' if number in image_cache or number in formula_cache else 'text_ready' if number in indexed else 'unindexed' if number in native else 'empty'
        labels = {'processing': '正在识别', 'failed': '识别未完成', 'pending': '图片/公式待核验', 'recognized': '已缓存识别结果',
                  'text_ready': '文字已提取', 'unindexed': '原文有文字，索引未包含', 'empty': '未提取到文字'}
        row = {'page': number, 'native_text': number in native, 'indexed_text': number in indexed,
               'image_candidate': number in images, 'formula_candidate': number in formulas,
               'recognized_images': number in image_cache, 'recognized_formulas': number in formula_cache,
               'recognition_methods': sorted(recognized.get(number, set())), 'status': status, 'label': labels[status],
               'retryable': bool(pending and not processing), 'last_attempt': last}
        if status == 'empty':
            row['detail'] = '未发现可提取文字或待识别的明显图片/公式对象；可能为空白页或矢量图，请查看原文确认。'
        elif pending:
            row['detail'] = '该页含图片或独立公式对象，相关内容尚未成功缓存；可单独识别。'
        else:
            row['detail'] = '提取结果可用于检索，仍请核对原文中的图表、符号和公式。'
        response['pages'].append(row)
    response['page_count'] = count
    response['summary'] = {name: [row['page'] for row in response['pages'] if predicate(row)] for name, predicate in {
        'native_text_pages': lambda row: row['native_text'], 'indexed_text_pages': lambda row: row['indexed_text'],
        'recognized_image_pages': lambda row: row['recognized_images'], 'recognized_formula_pages': lambda row: row['recognized_formulas'],
        'pending_pages': lambda row: row['status'] in {'pending', 'failed', 'processing'},
        'failed_pages': lambda row: row['status'] == 'failed', 'empty_pages': lambda row: row['status'] == 'empty',
    }.items()}
    return response


async def retry_document_pages(document: dict[str, Any], settings: dict[str, Any], pages: list[int],
                               on_page: Callable[[int], None] | None = None) -> dict[str, Any]:
    """Recognize only explicitly selected detected candidates, preserving existing chunks."""
    if not pages or any(type(page) is not int or page < 1 for page in pages):
        raise ValueError('请明确选择至少一个有效页码，不会自动改为全部来源。')
    requested = sorted(set(pages))
    async with _retry_locks.setdefault(document['id'], asyncio.Lock()):
        status = await run_in_threadpool(inspect_document, document)
        if status['error']:
            raise ValueError(status['error'])
        rows = {row['page']: row for row in status['pages']}
        if any(page not in rows or not rows[page]['retryable'] for page in requested):
            raise ValueError('所选页码没有待识别的图片或公式，请刷新解析状态后重新选择。')
        report = _read_report(document)
        now = datetime.now(UTC).isoformat()
        for page in requested:
            report['pages'][str(page)] = {'status': 'processing', 'updated_at': now}
        _write_report(document, report)
        _active_retries.add(document['id'])
        added: set[int] = set()
        unread: set[int] = set()
        try:
            image_pages = [page for page in requested if rows[page]['image_candidate'] and not rows[page]['recognized_images']]
            formula_pages = [page for page in requested if rows[page]['formula_candidate'] and not rows[page]['recognized_formulas']]
            for selected, recognizer in [(image_pages, ensure_image_text), (formula_pages, ensure_presentation_formulas)]:
                if selected:
                    result = await recognizer(document, settings, on_page=on_page, pages=selected)
                    added.update(result['added_pages'])
                    unread.update(result['unread_pages'])
        except asyncio.CancelledError:
            unread.update(requested)
            raise
        except (OSError, ValueError, RuntimeError, IndexError):
            unread.update(requested)
        finally:
            for page in requested:
                report['pages'][str(page)] = {'status': 'failed' if page in unread else 'complete',
                                             'updated_at': datetime.now(UTC).isoformat()}
            _active_retries.discard(document['id'])
            _write_report(document, report)
        refreshed = await run_in_threadpool(inspect_document, document)
        still_pending = set(refreshed['summary'].get('pending_pages', [])) & set(requested)
        # A later page may fail after an earlier page was committed. Re-read the
        # cache so successful pages are neither reported unread nor left out of
        # the caller's retrieval-index invalidation signal.
        for row in refreshed['pages']:
            old = rows.get(row['page'])
            if old and row['page'] in requested and any(row[field] and not old[field] for field in ('recognized_images', 'recognized_formulas')):
                added.add(row['page'])
        refreshed['retry'] = {'requested_pages': requested, 'added_pages': sorted(added), 'unread_pages': sorted(still_pending)}
        return refreshed
